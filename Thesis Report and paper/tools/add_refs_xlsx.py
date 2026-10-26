"""Add the four Bangladeshi counterfeit-detection papers to docs/literature/Review and paper links.xlsx (26 October 2026).
Facts from Crossref (DOIs) and the IJIRCST article page. Citation counts and abstracts are left empty: they were not
looked up, and are not guessed."""
import shutil
from copy import copy
from pathlib import Path

import openpyxl

X = Path(__file__).resolve().parents[2] / "docs" / "literature" / "Review and paper links.xlsx"
shutil.copy(X, X.with_name("Review and paper links_before_2026-10-26.xlsx"))
wb = openpyxl.load_workbook(X)
ws = wb["Review"]
last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value)
ROWS = [
    ("Evaluation of Deep Learning models on UV ink: a Fake Money detection scheme with RPN", "https://doi.org/10.1145/3542954.3542965",
     "March 2022", "Conference Paper",
     "A. A. Chowdhury, A. Das, D. Karmaker, K. K. S. Hoque (American International University-Bangladesh). ICCA 2022, ACM, pp. 67-74. "
     "Detects counterfeit Bangladeshi notes from the random colour flecks visible only under UV light; compares MobileNet, Inception, "
     "ResNet50/101 and Inception-ResNet; a region proposal network with CNN classification reaches 96.3% accuracy. Cited in the thesis "
     "(Literature review; UV fibre check)."),
    ("JaalTaka: A benchmark dataset for detecting counterfeit Bangladeshi banknotes", "https://doi.org/10.1016/j.dib.2025.112377",
     "February 2026", "Journal",
     "B. R. C. Piyas, F. J. Tisha, S. Rahman, S. Islam, B. R. C. Preenon. Data in Brief 64, 112377. 1,390 notes (802 genuine from government "
     "banks, 588 counterfeit from the Rapid Action Battalion), six views per note, worn, stained and repaired notes. Dataset: Mendeley Data "
     "10.17632/2m7wk5cy4c.2. The counterfeit benchmark used in the thesis."),
    ("Enhanced Counterfeit Detection of Bangladesh Currency through Convolutional Neural Networks: A Deep Learning Approach",
     "https://doi.org/10.55524/ijircst.2024.12.2.2", "2024", "Journal",
     "A. Pathak, A. Chakraborty, M. Rahaman, T. S. Rafa, U. Nayema. IJIRCST 12(2), pp. 10-20. CNNs trained on photographs of genuine and "
     "counterfeit Taka. Cited in the thesis (Literature review)."),
    ("Image-based approach for the detection of counterfeit banknotes of Bangladesh", "https://doi.org/10.1109/ICIEV.2016.7760162",
     "May 2016", "Conference Paper",
     "M. S. Uddin, P. P. Das, M. S. A. Roney. 5th International Conference on Informatics, Electronics and Vision (ICIEV), IEEE, pp. 1067-1072. "
     "Extracts watermark, micro-printing and hologram features from note images and classifies them with a support vector machine. "
     "Cited in the thesis (Literature review)."),
]
style_row = last                                   # copy the look of the last filled row
for k, (name, link, date, typ, summary) in enumerate(ROWS):
    r = last + 1 + k
    serial = (ws.cell(last, 1).value or last - 1) + 1 + k
    vals = {1: serial, 2: name, 3: link, 4: date, 5: None, 6: typ, 7: None, 8: summary}
    for c, v in vals.items():
        cell = ws.cell(r, c, v)
        src = ws.cell(style_row, c)
        if src.has_style:
            cell.font, cell.alignment, cell.border, cell.fill = copy(src.font), copy(src.alignment), copy(src.border), copy(src.fill)
            cell.number_format = src.number_format
    ws.cell(r, 3).hyperlink = link
wb.save(X)
print("added rows", last + 1, "to", last + len(ROWS))
